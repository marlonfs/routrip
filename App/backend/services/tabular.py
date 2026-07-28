import csv
import io

from openpyxl import load_workbook

from core.schemas import SpreadsheetSheet

MAX_ROWS = 1000
MAX_COLS = 30


def _fmt_cell(cell) -> str:
    if cell is None:
        return ""
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell).strip()


def _to_grid(rows, name: str) -> SpreadsheetSheet:
    grid: list[list[str]] = []
    total = 0
    for row in rows:
        total += 1
        if len(grid) >= MAX_ROWS:
            continue
        cells = [_fmt_cell(c) for c in row][:MAX_COLS]
        if any(cells):
            grid.append(cells)
    width = max((len(r) for r in grid), default=0)
    for row in grid:
        row.extend([""] * (width - len(row)))
    return SpreadsheetSheet(name=name, rows=grid, total_rows=total, truncated=total > MAX_ROWS)


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1", errors="replace")


def grid_from_csv(content: bytes) -> list[SpreadsheetSheet]:
    text = _decode(content)
    sample = text[:4096]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        # Planilhas brasileiras costumam usar ';'
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [_to_grid(reader, "Planilha")]


def grid_from_xlsx(content: bytes) -> list[SpreadsheetSheet]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        return [_to_grid(ws.iter_rows(values_only=True), ws.title) for ws in wb.worksheets]
    finally:
        wb.close()
